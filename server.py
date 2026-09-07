#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
移动端标签打印服务 - 内网使用，无安全措施
独立于 main.py 运行，仅依赖 openpyxl（已随 pandas 安装）
功能：将 PrintTools 桌面端核心功能移动端化
"""
import http.server
import json
import os
import socketserver
import sys
from datetime import datetime
from urllib.parse import urlparse

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    sys.exit(1)

PORT = 8000
BASE_DIR = os.path.dirname(os.path.abspath(__file__))


# ==================== 工具函数 ====================

def cell_str(c):
    """单元格值转字符串，处理 None 和浮点数"""
    if c is None:
        return ''
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c)


def try_num(s):
    """尝试将字符串转为数字，失败则返回原值（模拟 Excel OLE 自动转换）"""
    if not isinstance(s, str):
        return s
    if s == '':
        return None
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        pass
    return s


def read_config():
    try:
        with open(os.path.join(BASE_DIR, 'config.json'), 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def get_order_dir():
    cfg = read_config()
    d = cfg.get('orderExcelDirectory', BASE_DIR).replace('/', '\\')
    if not d.endswith('\\'):
        d += '\\'
    return d


def load_excel(path):
    """读取 xlsx，返回 (headers, rows) 或 None"""
    if not os.path.exists(path):
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not all_rows:
        return [], []
    headers = [cell_str(c) for c in all_rows[0]]
    rows = [[cell_str(c) for c in row] for row in all_rows[1:]]
    return headers, rows


def save_excel(filepath, headers, rows):
    """写入 xlsx，数值字符串自动转换"""
    wb = Workbook()
    ws = wb.active
    ws.append([try_num(h) for h in headers])
    for row in rows:
        ws.append([try_num(c) for c in row])
    wb.save(filepath)
    wb.close()


def load_exempt():
    path = os.path.join(BASE_DIR, 'exempt_config.json')
    if not os.path.exists(path):
        return {}
    try:
        with open(path, 'r', encoding='utf-8-sig') as f:
            return json.load(f)
    except Exception:
        return {}


def save_exempt_data(data):
    path = os.path.join(BASE_DIR, 'exempt_config.json')
    filtered = {k: v for k, v in data.items() if v and v > 0}
    with open(path, 'w', encoding='utf-8-sig') as f:
        json.dump(filtered, f, ensure_ascii=False, indent=2)


def make_save_path():
    """生成保存路径：orderExcelDirectory/yyyyMMdd_hhmmss手动打单.xlsx"""
    order_dir = get_order_dir()
    filename = datetime.now().strftime('%Y%m%d_%H%M%S') + '手动打单.xlsx'
    return os.path.join(order_dir, filename)


# ==================== HTML 页面 ====================

HTML_PAGE = r'''<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1.0,maximum-scale=1.0,user-scalable=no">
<title>标签打印</title>
<style>
*{margin:0;padding:0;box-sizing:border-box;-webkit-tap-highlight-color:transparent}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;background:#f0f2f5;color:#333;font-size:14px}
.tabs{display:flex;background:#fff;position:sticky;top:0;z-index:100;box-shadow:0 2px 4px rgba(0,0,0,.1)}
.tab{flex:1;padding:14px 0;text-align:center;font-size:16px;color:#666;border-bottom:3px solid transparent}
.tab.active{color:#1677ff;border-bottom-color:#1677ff;font-weight:600}
.content{padding:10px;padding-bottom:90px}
.search-bar{width:100%;padding:10px 14px;border:1px solid #ddd;border-radius:8px;font-size:16px;outline:none;margin-bottom:8px}
.chips{display:flex;flex-wrap:wrap;gap:6px;margin-bottom:10px}
.chip{padding:6px 14px;background:#fff;border:1px solid #ddd;border-radius:20px;font-size:13px;color:#666}
.chip.active{background:#1677ff;color:#fff;border-color:#1677ff}
.card{background:#fff;border-radius:8px;padding:12px;margin-bottom:8px;box-shadow:0 1px 3px rgba(0,0,0,.08)}
.card.marked{background:#f6ffed;border-left:3px solid #52c41a}
.card-title{font-size:16px;font-weight:600;margin-bottom:2px}
.card-sub{font-size:13px;color:#999;margin-bottom:6px}
.card-detail{font-size:12px;color:#bbb;margin-bottom:6px;line-height:1.4}
.controls{display:flex;gap:16px;align-items:center;flex-wrap:wrap}
.cg{display:flex;align-items:center;gap:6px}
.cg-label{font-size:13px;color:#888}
.cg-btn{width:34px;height:34px;border:none;border-radius:50%;font-size:18px;display:flex;align-items:center;justify-content:center;cursor:pointer;line-height:1}
.cg-btn.plus{background:#1677ff;color:#fff}
.cg-btn.minus{background:#e8e8e8;color:#333}
.cg-val{min-width:24px;text-align:center;font-size:16px;font-weight:600}
.save-bar{position:fixed;bottom:0;left:0;right:0;background:#fff;padding:10px;box-shadow:0 -2px 8px rgba(0,0,0,.1);display:flex;gap:10px;z-index:100}
.btn{flex:1;padding:14px;border:none;border-radius:8px;font-size:16px;font-weight:600;cursor:pointer}
.btn-primary{background:#1677ff;color:#fff}
.btn-outline{background:#fff;color:#1677ff;border:1px solid #1677ff}
.empty{text-align:center;padding:40px 20px;color:#999}
.toast{position:fixed;top:50%;left:50%;transform:translate(-50%,-50%);background:rgba(0,0,0,.8);color:#fff;padding:12px 24px;border-radius:8px;font-size:14px;z-index:9999;opacity:0;transition:opacity .3s;pointer-events:none}
.toast.show{opacity:1}
.refresh-btn{padding:6px 12px;background:#fff;border:1px solid #ddd;border-radius:6px;font-size:13px;color:#666;margin-bottom:8px}
</style>
</head>
<body>
<div class="tabs">
<div class="tab active" data-tab="manual" onclick="switchTab('manual')">手动打印</div>
<div class="tab" data-tab="latest" onclick="switchTab('latest')">最新打印</div>
</div>

<div id="tab-manual" class="content">
<input type="text" id="search" class="search-bar" placeholder="搜索商品..." oninput="onSearch()">
<div id="chips" class="chips"></div>
<div id="product-list"><div class="empty">加载中...</div></div>
</div>

<div id="tab-latest" class="content" style="display:none">
<button class="refresh-btn" onclick="loadLatest()">刷新</button>
<div id="latest-list"><div class="empty">加载中...</div></div>
</div>

<div id="save-bar-manual" class="save-bar">
<button class="btn btn-primary" onclick="saveManual()">生成打印表格</button>
</div>
<div id="save-bar-latest" class="save-bar" style="display:none">
<button class="btn btn-outline" onclick="saveLatestFirst()">补打第一条</button>
<button class="btn btn-primary" onclick="saveLatest()">最新打印补打</button>
</div>

<div id="toast" class="toast"></div>

<script>
// ===== 状态 =====
let gHeaders=[], gRows=[], gConfig={}, exemptCfg={};
let quantities={}, exempts={};
let lHeaders=[], lRows=[];
let lMarkStart=-1;
let activeTab='manual', activeChip=null, searchText='';
let specIdIdx=-1, latestLoaded=false;

// ===== 工具 =====
function showToast(msg){
  const t=document.getElementById('toast');
  t.textContent=msg; t.classList.add('show');
  setTimeout(()=>t.classList.remove('show'),2000);
}
function colIdx(headers,name){
  return headers.findIndex(h=>h&&h.trim()===name);
}
async function api(path,opts){
  const r=await fetch(path,opts);
  return r.json();
}

// ===== 初始化 =====
async function init(){
  try{
    const [g,c,e]=await Promise.all([
      api('/api/goods'), api('/api/goods_config'), api('/api/exempt')
    ]);
    if(g.error){document.getElementById('product-list').innerHTML='<div class="empty">goods.xlsx 未找到</div>';return;}
    gHeaders=g.headers; gRows=g.rows; gConfig=c; exemptCfg=e;
    specIdIdx=colIdx(gHeaders,'规格ID');
    // 从免打配置初始化
    gRows.forEach((row,i)=>{
      if(specIdIdx>=0){
        const sid=row[specIdIdx];
        if(sid&&exemptCfg[sid]) exempts[i]=exemptCfg[sid];
      }
    });
    renderChips();
    renderProducts();
  }catch(e){showToast('加载失败: '+e.message);}
}

function renderChips(){
  const c=document.getElementById('chips');
  c.innerHTML='';
  Object.entries(gConfig).forEach(([name,py])=>{
    const d=document.createElement('div');
    d.className='chip'; d.textContent=name;
    d.onclick=()=>toggleChip(name,py);
    c.appendChild(d);
  });
}

function toggleChip(name,py){
  if(activeChip===py){
    activeChip=null;
    document.getElementById('search').value='';
    searchText='';
  }else{
    activeChip=py;
    document.getElementById('search').value=py;
    searchText=py;
  }
  document.querySelectorAll('.chip').forEach(c=>{
    c.classList.toggle('active',c.textContent===name&&activeChip===py);
  });
  renderProducts();
}

function onSearch(){
  searchText=document.getElementById('search').value;
  activeChip=null;
  document.querySelectorAll('.chip').forEach(c=>c.classList.remove('active'));
  renderProducts();
}

function renderProducts(){
  const c=document.getElementById('product-list');
  c.innerHTML='';
  const ni=colIdx(gHeaders,'成分');
  const si=colIdx(gHeaders,'规格');
  const wi=colIdx(gHeaders,'净重');
  let count=0;
  gRows.forEach((row,i)=>{
    // 筛选
    if(searchText){
      const m=row.some(v=>v&&v.includes(searchText));
      if(!m) return;
    }
    const card=document.createElement('div');
    card.className='card';
    const name=ni>=0?row[ni]:(row[0]||'');
    const spec=si>=0?row[si]:'';
    const weight=wi>=0?row[wi]:'';
    let h='<div class="card-title">'+name+'</div>';
    if(spec||weight) h+='<div class="card-sub">'+spec+' '+weight+'</div>';
    // 其他字段
    const details=[];
    gHeaders.forEach((hd,j)=>{
      if(j!==ni&&j!==si&&j!==wi&&j!==specIdIdx&&row[j]){
        details.push(hd+': '+row[j]);
      }
    });
    if(details.length) h+='<div class="card-detail">'+details.join(' | ')+'</div>';
    // 控件
    const q=quantities[i]||0, ex=exempts[i]||0;
    h+='<div class="controls">'+
      '<div class="cg"><span class="cg-label">数量</span>'+
      '<button class="cg-btn minus" onclick="changeQty('+i+',-1)">-</button>'+
      '<span class="cg-val">'+q+'</span>'+
      '<button class="cg-btn plus" onclick="changeQty('+i+',1)">+</button></div>'+
      '<div class="cg"><span class="cg-label">免打</span>'+
      '<button class="cg-btn minus" onclick="changeExempt('+i+',-1)">-</button>'+
      '<span class="cg-val">'+ex+'</span>'+
      '<button class="cg-btn plus" onclick="changeExempt('+i+',1)">+</button></div>'+
      '</div>';
    card.innerHTML=h;
    c.appendChild(card);
    count++;
  });
  if(count===0) c.innerHTML='<div class="empty">没有匹配的商品</div>';
}

function changeQty(idx,delta){
  const v=Math.max(0,(quantities[idx]||0)+delta);
  quantities[idx]=v;
  renderProducts();
}

async function changeExempt(idx,delta){
  const v=Math.max(0,(exempts[idx]||0)+delta);
  exempts[idx]=v;
  if(specIdIdx>=0){
    const sid=gRows[idx][specIdIdx];
    if(sid){
      try{
        await api('/api/exempt',{
          method:'POST',
          headers:{'Content-Type':'application/json'},
          body:JSON.stringify({[sid]:v})
        });
      }catch(e){showToast('保存免打失败');}
    }
  }
  renderProducts();
}

async function saveManual(){
  const headers=['数量','免打',...gHeaders];
  const rows=[];
  gRows.forEach((row,i)=>{
    const q=quantities[i]||0;
    if(q>0){
      const ex=exempts[i]||'';
      rows.push([String(q),String(ex),...row]);
    }
  });
  if(!rows.length){showToast('没有设置数量的商品');return;}
  try{
    const r=await api('/api/save_manual',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({headers,rows})
    });
    if(r.ok){
      showToast('已生成: '+r.file);
      quantities={};
      renderProducts();
    }else showToast('保存失败: '+(r.error||''));
  }catch(e){showToast('保存失败: '+e.message);}
}

// ===== 最新打印 =====
async function loadLatest(){
  try{
    const r=await api('/api/latest_print');
    if(r.error){
      document.getElementById('latest-list').innerHTML='<div class="empty">暂无最新打印文件</div>';
      lHeaders=[]; lRows=[];
      return;
    }
    lHeaders=r.headers; lRows=r.rows; lMarkStart=-1;
    renderLatest();
  }catch(e){showToast('加载失败: '+e.message);}
}

function renderLatest(){
  const c=document.getElementById('latest-list');
  c.innerHTML='';
  if(!lRows.length){c.innerHTML='<div class="empty">暂无数据</div>';return;}
  const ni=colIdx(lHeaders,'成分');
  const si=colIdx(lHeaders,'规格');
  const wi=colIdx(lHeaders,'净重');
  lRows.forEach((row,i)=>{
    const card=document.createElement('div');
    card.className='card';
    if(lMarkStart>=0&&i>=lMarkStart) card.classList.add('marked');
    const name=ni>=0?row[ni]:(row[0]||'');
    const spec=si>=0?row[si]:'';
    const weight=wi>=0?row[wi]:'';
    let h='<div class="card-title">'+name+'</div>';
    if(spec||weight) h+='<div class="card-sub">'+spec+' '+weight+'</div>';
    const details=[];
    lHeaders.forEach((hd,j)=>{
      if(j!==ni&&j!==si&&j!==wi&&row[j]){
        details.push(hd+': '+row[j]);
      }
    });
    if(details.length) h+='<div class="card-detail">'+details.join(' | ')+'</div>';
    h+='<div class="card-detail" style="color:#1677ff">点击标记此行及以下</div>';
    card.innerHTML=h;
    card.onclick=()=>markLatest(i);
    c.appendChild(card);
  });
}

function markLatest(idx){
  lMarkStart=(lMarkStart===idx)?-1:idx;
  renderLatest();
}

async function saveLatest(){
  if(lMarkStart<0){showToast('没有标记的行');return;}
  const headers=['数量','免打',...lHeaders];
  const rows=[];
  for(let i=lMarkStart;i<lRows.length;i++){
    rows.push(['1','',...lRows[i]]);
  }
  try{
    const r=await api('/api/save_latest',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({headers,rows})
    });
    if(r.ok) showToast('已生成: '+r.file);
    else showToast('保存失败: '+(r.error||''));
  }catch(e){showToast('保存失败: '+e.message);}
}

async function saveLatestFirst(){
  if(!lRows.length){showToast('暂无数据');return;}
  const headers=['数量','免打',...lHeaders];
  const row=['1','',...lRows[0]];
  try{
    const r=await api('/api/save_latest_first',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({headers,row})
    });
    if(r.ok) showToast('已生成: '+r.file);
    else showToast('保存失败: '+(r.error||''));
  }catch(e){showToast('保存失败: '+e.message);}
}

// ===== Tab 切换 =====
function switchTab(tab){
  activeTab=tab;
  document.querySelectorAll('.tab').forEach((t)=>{
    t.classList.toggle('active',t.dataset.tab===tab);
  });
  document.getElementById('tab-manual').style.display=tab==='manual'?'':'none';
  document.getElementById('tab-latest').style.display=tab==='latest'?'':'none';
  document.getElementById('save-bar-manual').style.display=tab==='manual'?'':'none';
  document.getElementById('save-bar-latest').style.display=tab==='latest'?'':'none';
  if(tab==='latest') loadLatest();
}

init();
</script>
</body>
</html>'''


# ==================== 请求处理 ====================

class Handler(http.server.BaseHTTPRequestHandler):
    protocol_version = 'HTTP/1.1'

    def log_message(self, *args):
        pass  # 静默

    def _send_json(self, data, code=200):
        body = json.dumps(data, ensure_ascii=False).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _send_html(self, html):
        body = html.encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _read_body(self):
        length = int(self.headers.get('Content-Length', 0))
        if length == 0:
            return {}
        return json.loads(self.rfile.read(length).decode('utf-8'))

    def do_GET(self):
        path = urlparse(self.path).path
        if path in ('/', '/index.html'):
            self._send_html(HTML_PAGE)
        elif path == '/api/goods':
            result = load_excel(os.path.join(BASE_DIR, 'goods.xlsx'))
            if result is None:
                self._send_json({'error': 'goods.xlsx not found'})
            else:
                self._send_json({'headers': result[0], 'rows': result[1]})
        elif path == '/api/goods_config':
            try:
                with open(os.path.join(BASE_DIR, 'goods_config.json'), 'r', encoding='utf-8') as f:
                    self._send_json(json.load(f))
            except Exception:
                self._send_json({})
        elif path == '/api/exempt':
            self._send_json(load_exempt())
        elif path == '/api/latest_print':
            result = load_excel(os.path.join(BASE_DIR, 'latest_print_file.xlsx'))
            if result is None:
                self._send_json({'error': 'no file'})
            else:
                self._send_json({'headers': result[0], 'rows': result[1]})
        else:
            self.send_response(404)
            self.end_headers()

    def do_POST(self):
        path = urlparse(self.path).path
        if path == '/api/exempt':
            data = self._read_body()
            current = load_exempt()
            for k, v in data.items():
                if v and v > 0:
                    current[k] = v
                else:
                    current.pop(k, None)
            save_exempt_data(current)
            self._send_json({'ok': True})
        elif path in ('/api/save_manual', '/api/save_latest', '/api/save_latest_first'):
            data = self._read_body()
            headers = data.get('headers', [])
            if path == '/api/save_latest_first':
                rows = [data.get('row', [])]
            else:
                rows = data.get('rows', [])
            if not rows or not headers:
                self._send_json({'ok': False, 'error': '没有数据'})
                return
            filepath = make_save_path()
            try:
                save_excel(filepath, headers, rows)
                self._send_json({'ok': True, 'file': os.path.basename(filepath)})
            except Exception as e:
                self._send_json({'ok': False, 'error': str(e)})
        else:
            self.send_response(404)
            self.end_headers()


# ==================== 主程序 ====================

class ThreadingHTTPServer(socketserver.ThreadingMixIn, http.server.HTTPServer):
    daemon_threads = True
    allow_reuse_address = True


def main():
    try:
        server = ThreadingHTTPServer(('0.0.0.0', PORT), Handler)
        server.serve_forever()
    except OSError:
        pass  # 端口被占用，静默退出（已有实例在运行）
    except KeyboardInterrupt:
        pass


if __name__ == '__main__':
    main()
